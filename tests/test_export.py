import csv
import json
from pathlib import Path

import openpyxl

from scraper_app.exporting.excel import export_csv, export_excel, export_json


def test_excel_csv_json_exports(tmp_path: Path) -> None:
    rows = [
        {"title": "A", "price": 12.5, "url": "https://example.test/item/1"},
        {"title": "=not-a-formula", "price": 9, "url": "not a url"},
    ]
    excel = export_excel(iter(rows), tmp_path / "results.xlsx", "https://example.test", "Example")
    csv_path = export_csv(rows, tmp_path / "results.csv")
    json_path = export_json(rows, tmp_path / "results.json")

    workbook = openpyxl.load_workbook(excel, data_only=False)
    sheet = workbook["Data 1"]
    assert sheet.freeze_panes == "A2"
    assert sheet["A1"].value == "title"
    assert sheet["B2"].value == 12.5
    assert sheet["A3"].value == "=not-a-formula"
    assert sheet["C2"].hyperlink.target == "https://example.test/item/1"
    assert workbook["Ringkasan"]["B5"].value == 2
    with csv_path.open(encoding="utf-8-sig", newline="") as stream:
        csv_rows = list(csv.DictReader(stream))
        assert csv_rows[0]["title"] == "A"
        assert csv_rows[1]["title"] == "'=not-a-formula"
    assert json.loads(json_path.read_text(encoding="utf-8")) == rows


def test_excel_splits_streamed_records_across_sheets(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr("scraper_app.exporting.excel.EXCEL_DATA_ROWS", 2)
    rows = ({"value": value} for value in ("A", "B", "C"))
    path = export_excel(rows, tmp_path / "split.xlsx", "https://example.test", "Example")

    workbook = openpyxl.load_workbook(path, read_only=True)
    assert workbook.sheetnames == ["Data 1", "Data 2", "Ringkasan"]
    assert list(workbook["Data 1"].values) == [("value",), ("A",), ("B",)]
    assert list(workbook["Data 2"].values) == [("value",), ("C",)]
    assert workbook["Ringkasan"]["B5"].value == 3
